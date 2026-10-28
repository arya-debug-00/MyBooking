import openpyxl
import os

HEADERS = [
    "username",
    "password",
    "expected_result",
    "Kollinumber",
    "ExcelSenderName",
    "ExcelReceiverName",
    "ExcelconsignmentNo",
    "ExcelSenderAddress1",
    "ExcelReceiverAddress1",
    "ExcelSenderPostalCode",
    "ExcelReceiverPostalCode",
    "ExcelEvent"
]

# def get_test_data(relative_path):
#     base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#     full_path = os.path.join(base_dir, relative_path)

#     if not os.path.exists(full_path):
#         raise FileNotFoundError(f"Excel file not found: {full_path}")

#     workbook = openpyxl.load_workbook(full_path)
#     sheet = workbook.active
#     headers = [cell.value for cell in sheet[1]]
#     data = []

#     for row in sheet.iter_rows(min_row=2, values_only=True):
#         row_dict = {header: (row[i] if i < len(row) else "") for i, header in enumerate(headers)}
#         data.append(row_dict)

#     return data



def get_test_data(relative_path, sheet_name=None):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    full_path = os.path.join(base_dir, relative_path)

    if not os.path.exists(full_path):
        raise FileNotFoundError(f"Excel file not found: {full_path}")

    workbook = openpyxl.load_workbook(full_path)

    if sheet_name:
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active

    headers = [str(cell.value).strip() for cell in sheet[1]]

    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_dict = {
            header: str(row[i]).strip() if i < len(row) and row[i] is not None else ""
            for i, header in enumerate(headers)
        }
        data.append(row_dict)

    return data
